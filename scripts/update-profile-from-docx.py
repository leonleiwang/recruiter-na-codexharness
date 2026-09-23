#!/usr/bin/env python3
"""用新版画像 DOCX 更新画像库中的同名条目（保留 profile_id 等既有字段）。

DOCX 结构：单列表格，依次是 岗位差异化要求 / 业务特征挑战 / 【关键能力要求】 /
【关键经验要求】 / 【管理经验】 / 【其他特质要求】 / 版本更新时间。

用法：python scripts/update-profile-from-docx.py --docx a.docx --docx b.docx
"""
from __future__ import annotations

import argparse
import json
import re
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
LIBRARY = ROOT / "data" / "generated" / "profiles" / "profile_library.json"


def read_blocks(path: Path) -> list[str]:
    import docx
    doc = docx.Document(str(path))
    blocks: list[str] = []
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                text = cell.text.strip()
                if text and text not in blocks:
                    blocks.append(text)
    for para in doc.paragraphs:
        text = para.text.strip()
        if text and text not in blocks:
            blocks.append(text)
    return blocks


def find_block(blocks: list[str], marker: str) -> str:
    for block in blocks:
        if marker in block:
            return block
    return ""


def parse_items(block: str) -> list[dict[str, str]]:
    body = block.split("】：", 1)[-1] if "】：" in block else block
    items: list[dict[str, str]] = []
    for raw in body.splitlines():
        line = raw.strip()
        if not line:
            continue
        m = re.match(r"^\s*\d+\s*([^：:]{2,30})[：:]\s*(.+)$", line)
        if not m:
            continue
        name = m.group(1).strip()
        definition = m.group(2).strip()
        if name and definition:
            items.append({"name": name, "def": definition})
    return items


def parse_traits(block: str) -> list[str]:
    body = block.split("】", 1)[-1] if "】" in block else block
    out: list[str] = []
    for part in re.split(r"[；;]", body):
        piece = part.strip()
        if not piece:
            continue
        name = re.split(r"[：:]", piece, maxsplit=1)[0].strip()
        if name and name not in out:
            out.append(name)
    return out


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("--docx", action="append", required=True)
    args = ap.parse_args()

    library = json.loads(LIBRARY.read_text(encoding="utf-8"))
    profiles = library.get("profiles") or []
    changed: list[str] = []

    for raw_path in args.docx:
        path = Path(raw_path)
        blocks = read_blocks(path)
        known = {p.get("name") for p in profiles}
        title = next((b for b in blocks if b in known), "")
        if not title:
            title = path.stem.split("-", 1)[-1] if "-" in path.stem else path.stem
        entry = next((p for p in profiles if p.get("name") == title), None)
        if entry is None:
            print("跳过（画像库无同名条目）:", title)
            continue

        abilities = parse_items(find_block(blocks, "关键能力要求"))
        experiences = parse_items(find_block(blocks, "关键经验要求"))
        traits = parse_traits(find_block(blocks, "其他特质要求"))
        challenge = find_block(blocks, "业务特征")
        management = find_block(blocks, "管理经验")
        version = re.search(r"版本更新时间[：:]\s*([0-9.]+)", chr(10).join(blocks))

        before = (
            len(entry.get("key_abilities") or []),
            len(entry.get("key_experiences") or []),
            len(entry.get("key_traits") or []),
        )
        if abilities:
            entry["key_abilities"] = abilities
        if experiences:
            entry["key_experiences"] = experiences
        if traits:
            entry["key_traits"] = traits
        if challenge:
            entry["business_challenge"] = challenge.split("：", 1)[-1].strip()
        if management:
            entry["management"] = management.replace(chr(10), " ").strip()
        if version:
            entry["source_version_date"] = version.group(1)
        entry["source_file"] = path.name
        after = (
            len(entry["key_abilities"]),
            len(entry["key_experiences"]),
            len(entry["key_traits"]),
        )
        changed.append("%s: 能力/经验/特质 %s -> %s | 版本 %s" % (
            entry.get("name"), before, after, entry.get("source_version_date")))

    LIBRARY.write_text(json.dumps(library, ensure_ascii=False, indent=2), encoding="utf-8")
    for line in changed:
        print(line)
    print("已更新画像库:", LIBRARY)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
