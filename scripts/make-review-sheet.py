"""Build the human-review sheet (one row per candidate) from an audit.json.

The recruiter fills the「人工结论」column (下拉：推荐 / 待核实 / 不推荐) and any remarks,
then the sheet is read back by ``scripts/eval-matching.py --review-xlsx`` to compute
agreement with the rule-based conclusion.

Usage (from the repo root, Windows):

    python scripts\\make-review-sheet.py ^
        --audit workspace\\runs\\matching-accuracy\\audit.json ^
        --out workspace\\runs\\matching-accuracy\\人工校验表.xlsx
"""

from __future__ import annotations

import argparse
import json
from pathlib import Path
import re
import sys

try:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.datavalidation import DataValidation
except ImportError:  # pragma: no cover - openpyxl is a bundled dependency
    print("openpyxl is required: pip install openpyxl", file=sys.stderr)
    raise SystemExit(2)

STATUS_MARK = {
    "matched": "命中",
    "weak": "待验证",
    "not_found": "待验证（未提及）",
    "contradicted": "未命中（反证）",
    "interview_pending": "待面试核",
}
HARD_FILTERS = (("education", "学历核验"), ("certification", "证书核验"), ("location", "工作地点核验"))
FILL = {
    "命中": PatternFill("solid", fgColor="E3F1E3"),
    "待验证": PatternFill("solid", fgColor="FFF3D6"),
    "待验证（未提及）": PatternFill("solid", fgColor="FFF3D6"),
    "未命中（反证）": PatternFill("solid", fgColor="F2F2F2"),
}
THIN = Side(style="thin", color="D9D9D9")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
HEADER_FONT = Font(color="FFFFFF", bold=True)


def _resume_name(file_name: str) -> str:
    return Path(file_name).stem.split("_")[0]


def _hard_status(hard_filters: list[dict], key: str) -> str:
    for hint in hard_filters:
        if hint.get("filter") == key:
            return str(hint.get("status") or "")
    return ""


def _hard_notes(hard_filters: list[dict]) -> str:
    lines = []
    for hint in hard_filters:
        note = hint.get("note") or hint.get("status") or ""
        evidence = re.sub(r"\s+", " ", str(hint.get("evidence") or "")).strip()
        line = f"{hint.get('filter')}：{note}".strip()
        if evidence:
            line += f"\n依据：{_clip(evidence, 60)}"
        lines.append(line)
    return "\n".join(lines)


MODE_LABEL = {"resume_evidence": "简历证据", "interview_verify": "面试核实"}


def _clip(text: str, limit: int = 60) -> str:
    text = re.sub(r"\s+", " ", str(text or "")).strip()
    return text if len(text) <= limit else text[:limit] + "…"


def _cell_text(item: dict) -> str:
    """Tag cell: the mark on the first line, then the evidence it was decided on.

    填表人不用再去 audit.md 里翻证据：命中写命中的原句/关键词，待验证写线索或「未提及」，
    反证写否定句。
    """
    mark = STATUS_MARK.get(item["state"], item["state"])
    evidence = _clip(item.get("evidence") or "", 70)
    words = [
        str(word)
        for word in (item.get("strong_keywords") or item.get("matched_keywords") or [])
        if str(word).strip()
    ]
    parts: list[str] = []
    if evidence:
        parts.append("依据：" + evidence)
    elif mark.startswith("待验证（未提及）"):
        parts.append("依据：简历未提及")
    if words:
        parts.append("词：" + "、".join(words[:4]) + ("…" if len(words) > 4 else ""))
    return "\n".join([mark, *parts])


def _tag_header(item: dict) -> str:
    """Header cell: tag name + level + 判定口径 + 关键词（填表人不用翻文档）。"""
    lines = [f"{item['name']}（{item['level']}）"]
    mode = MODE_LABEL.get(item.get("judge_mode") or "", item.get("judge_mode") or "简历证据")
    lines.append(f"判据：{mode}")
    strong = [str(word) for word in (item.get("judge_strong_keywords") or []) if str(word).strip()]
    words = strong or [str(word) for word in (item.get("judge_keywords") or []) if str(word).strip()]
    if words:
        detail = "、".join(words[:6]) + ("…" if len(words) > 6 else "")
        lines.append(("强词：" if strong else "词：") + detail)
    return "\n".join(lines)


def build_sheet(rows: list[dict]) -> Workbook:
    wb = Workbook()
    ws = wb.active
    ws.title = "人工校验"

    items = rows[0]["items"] if rows else []
    headers = ["序号", "姓名", "简历文件"]
    headers += [label for _, label in HARD_FILTERS]
    headers += ["硬条件说明"]
    headers += [_tag_header(item) for item in items]
    headers += ["命中/适用", "命中率", "系统结论", "人工结论", "人工意见", "备注"]
    ws.append(headers)
    for column, _ in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=column)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = BORDER

    for index, row in enumerate(rows, start=1):
        line = [index, _resume_name(row["file"]), row["file"]]
        line += [_hard_status(row["hard_filters"], key) for key, _ in HARD_FILTERS]
        line.append(_hard_notes(row["hard_filters"]))
        line += [_cell_text(item) for item in row["items"]]
        line += [
            f"{row['tag_hit_count']}/{row['tag_applied_count']}",
            row["hit_percent"] / 100,
            row["label"],
            "",
            "",
            "",
        ]
        ws.append(line)
        for column in range(1, len(headers) + 1):
            ws.cell(row=index + 1, column=column).border = BORDER
        for column in range(1, len(headers) + 1):
            cell = ws.cell(row=index + 1, column=column)
            cell.alignment = Alignment(vertical="center", wrap_text=column in (7, len(headers) - 1, len(headers)))
        for offset, item in enumerate(items):
            cell = ws.cell(row=index + 1, column=4 + len(HARD_FILTERS) + 1 + offset)
            mark = STATUS_MARK.get(item["state"], item["state"])
            if mark in FILL:
                cell.fill = FILL[mark]
            cell.alignment = Alignment(vertical="top", wrap_text=True)
        ws.row_dimensions[index + 1].height = 78
        percent_cell = ws.cell(row=index + 1, column=len(headers) - 4)
        percent_cell.number_format = "0.0%"

    review_column = len(headers) - 2
    letter = get_column_letter(review_column)
    validation = DataValidation(
        type="list", formula1='"推荐,待核实,不推荐"', allow_blank=True, showDropDown=False
    )
    validation.error = "请选择：推荐 / 待核实 / 不推荐"
    validation.errorTitle = "人工结论"
    ws.add_data_validation(validation)
    validation.add(f"{letter}2:{letter}{max(len(rows) + 1, 2)}")

    widths = {"A": 5, "B": 9, "C": 24, "D": 9, "E": 9, "F": 12, "G": 46}
    for column, width in widths.items():
        ws.column_dimensions[column].width = width
    for offset in range(len(items)):
        ws.column_dimensions[get_column_letter(4 + len(HARD_FILTERS) + 1 + offset)].width = 30
    tail = len(headers)
    for offset in range(6):
        ws.column_dimensions[get_column_letter(tail - 5 + offset)].width = 12 if offset < 3 else 26
    ws.column_dimensions[get_column_letter(review_column)].width = 13
    ws.column_dimensions[get_column_letter(review_column + 1)].width = 30
    ws.column_dimensions[get_column_letter(review_column + 2)].width = 20
    ws.freeze_panes = "C2"
    ws.row_dimensions[1].height = 110
    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{len(rows) + 1}"

    guide = wb.create_sheet("填表说明")
    guide.column_dimensions["A"].width = 110
    guide_lines = [
        "填表说明（人工校验用）",
        "",
        "1. 每一行是一位候选人；8 个标签列是系统按 JD 拆出的评分标签。",
        "   每个标签格第一行是系统判定（命中 / 待验证 / 未命中（反证）），下面几行是判定依据：",
        "   「依据：」是简历原文证据，「词：」是命中的关键词——你只需核对依据是否成立。",
        "2. 请在「人工结论」列用下拉选择你的判断：推荐 / 待核实 / 不推荐；可在「人工意见」写理由，",
        "   「备注」写证件、地点、经验等需要补充核实的点。",
        "3. 判定口径（系统）：",
        "   - 命中：简历有明确证据（岗位强关键词，或同维度 ≥2 个有区分度的关键词）；",
        "   - 待验证（有线索）：只有 1 条弱线索，或简历写「备考/待考/未取得」等语境；",
        "   - 待验证（未提及）：简历没写；没写不等于不会，仍计入命中率分母，面试时逐项确认；",
        "   - 未命中（反证）：简历明确写了没有/不符合（如明确无相关经历、专业不符），才排除。",
        "   - 硬条件（学历/证书/工作地点）单独核验：不满足即不推荐，简历没写明记待核实。",
        "4. 命中率 = 命中标签数 / 适用标签数（只算能力/经验，不含特质），仅供辅助排序，不是录用评分。",
        "5. 填完把 xlsx 发回即可；eval-matching.py --review-xlsx 会统计人工结论与系统结论的一致率。",
        "6. 一致率对照口径：高匹配/较高匹配 → 推荐，中等匹配 → 待核实；命中率不足四成的偏低匹配也记待核实。",
        "   所以「人工结论」请按三档填；若你的判断和系统不同，直接在「人工意见」里写理由即可。",
    ]
    for line in guide_lines:
        guide.append([line])
    guide["A1"].font = Font(bold=True, size=13)

    return wb


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description="Build the human-review xlsx from audit.json")
    parser.add_argument("--audit", required=True)
    parser.add_argument("--out", required=True)
    args = parser.parse_args(argv)

    payload = json.loads(Path(args.audit).read_text(encoding="utf-8"))
    rows = payload.get("rows") or []
    if not rows:
        parser.error(f"no rows in {args.audit}")
    out = Path(args.out)
    out.parent.mkdir(parents=True, exist_ok=True)
    build_sheet(rows).save(out)
    print(json.dumps({"out": str(out), "rows": len(rows)}, ensure_ascii=False))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
