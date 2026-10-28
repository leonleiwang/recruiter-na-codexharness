#!/usr/bin/env python
"""Build a tag-spec sheet: every scoring tag + its matching basis, for human review.

The sheet is meant to be filled in by the recruiter to correct the tag definitions
(keywords / judge basis), not the candidate verdicts. One tag per row, plus the
current per-resume states so a wrong cell points at the keyword to fix.

Usage:
    python scripts/make-tag-spec-sheet.py --audit workspace/runs/matching-accuracy/audit.json \
        --out workspace/runs/matching-accuracy/标签口径表.xlsx
"""

from __future__ import annotations

import argparse
import json
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
HEADER_FONT = Font(color="FFFFFF", bold=True, size=10)
TITLE_FONT = Font(bold=True, size=13, color="1F4E79")
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
STATUS_FILL = {
    "matched": PatternFill("solid", fgColor="E3F1E3"),
    "weak": PatternFill("solid", fgColor="FFF3D6"),
    "not_found": PatternFill("solid", fgColor="F2F2F2"),
}
STATUS_LABEL = {
    "matched": "命中",
    "weak": "待验证（有线索）",
    "not_found": "待验证（未提及）",
    "contradicted": "未命中（反证）",
}
DIMENSION_LABEL = {"capability": "能力", "experience": "经验", "trait": "特质"}
MODE_LABEL = {"resume_evidence": "简历证据", "interview_verify": "面试核实"}

COLUMNS = [
    ("序号", 6),
    ("标签名", 20),
    ("维度", 7),
    ("级别", 7),
    ("判据方式", 10),
    ("命中依据（强词，命中即算匹配）", 46),
    ("弱线索词（≥2 个有区分度的词才匹配）", 34),
    ("判定规则", 26),
    ("本批命中", 9),
    ("示例证据（简历原文）", 52),
    ("人工修正（可直接写：删 XX / 补 XX / 这条不算）", 30),
    ("备注", 18),
]
RESUME_WIDTH = 11


def _short_name(path: str) -> str:
    name = Path(path).stem
    return name.split("_")[0]


def _load(audit_path: Path) -> dict:
    return json.loads(audit_path.read_text(encoding="utf-8"))


def _tag_rows(audit: dict) -> list[dict]:
    """Aggregate per-tag definitions plus per-resume states from the audit rows."""
    tags: dict[str, dict] = {}
    for row in audit.get("rows") or []:
        for item in row.get("items") or []:
            rid = str(item.get("requirement_id") or "")
            if not rid:
                continue
            tag = tags.setdefault(rid, {
                "id": rid,
                "name": item.get("name") or "",
                "dimension": item.get("dimension") or "",
                "level": item.get("level") or "",
                "judge_mode": item.get("judge_mode") or "",
                "keywords": list(item.get("judge_keywords") or []),
                "strong": list(item.get("judge_strong_keywords") or []),
                "states": {},
                "evidence": [],
                "hits": 0,
            })
            tag["states"][row["file"]] = item.get("state") or "not_found"
            if (item.get("state") or "") == "matched":
                tag["hits"] += 1
            snippet = str(item.get("evidence") or "").strip()
            if snippet and snippet not in tag["evidence"]:
                tag["evidence"].append(snippet)
    return list(tags.values())


def _rule_text(tag: dict) -> str:
    if tag["level"] == "优先":
        return "命中=匹配；单条弱线索=待验证；没写=待验证（未提及，仍计分母）"
    return "强词命中=匹配；≥2 条弱词=匹配；1 条弱词=待验证；没写=待验证（未提及，仍计分母）"


def _style_header(ws, ncols: int, height: int) -> None:
    for col in range(1, ncols + 1):
        cell = ws.cell(row=1, column=col)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = BORDER
    ws.row_dimensions[1].height = height


def _write_spec_sheet(ws, tags: list[dict], resume_files: list[str]) -> None:
    headers = [name for name, _ in COLUMNS] + [_short_name(f) for f in resume_files]
    widths = [width for _, width in COLUMNS] + [RESUME_WIDTH] * len(resume_files)
    ws.append(headers)
    for index, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(index)].width = width
    _style_header(ws, len(headers), 60)

    for seq, tag in enumerate(tags, start=1):
        evidence = " ｜ ".join(tag["evidence"][:2])
        if len(evidence) > 220:
            evidence = evidence[:220] + "…"
        values = [
            seq,
            tag["name"],
            DIMENSION_LABEL.get(tag["dimension"], tag["dimension"]),
            tag["level"],
            MODE_LABEL.get(tag["judge_mode"], tag["judge_mode"]),
            "、".join(tag["strong"]) or "（无强词）",
            "、".join(k for k in tag["keywords"] if k not in tag["strong"]) or "（无）",
            _rule_text(tag),
            f"{tag['hits']}/{len(resume_files)}",
            evidence,
            "",
            "",
        ]
        ws.append(values)
        row = ws.max_row
        for col in range(1, len(headers) + 1):
            cell = ws.cell(row=row, column=col)
            cell.border = BORDER
            cell.alignment = Alignment(vertical="top", wrap_text=True)
        ws.row_dimensions[row].height = 64
        ws.cell(row=row, column=1).alignment = Alignment(horizontal="center", vertical="top")
        for index, resume in enumerate(resume_files):
            cell = ws.cell(row=row, column=len(COLUMNS) + 1 + index)
            state = tag["states"].get(resume) or "not_found"
            cell.value = STATUS_LABEL.get(state, state)
            cell.fill = STATUS_FILL.get(state, STATUS_FILL["not_found"])
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = BORDER
        for col in (11, 12):
            ws.cell(row=row, column=col).fill = PatternFill("solid", fgColor="FFF9E6")

    ws.freeze_panes = "C2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{ws.max_row}"


def _write_notes_sheet(ws) -> None:
    ws.column_dimensions["A"].width = 120
    lines = [
        ("标签口径与依据核对说明", TITLE_FONT),
        ("", None),
        ("1. 这张表一行一个标签：左边是标签定义，右边 10 列是这 10 份简历当前的判定结果。", None),
        ("2. 要改的是「依据」列：在「命中依据（强词）」里直接增删词，用顿号分隔；", None),
        ("   「弱线索词」是只算线索、要 ≥2 个才构成命中的词。", None),
        ("3. 判定口径：", None),
        ("   - 命中（match）：简历里有强词，或同维度 ≥2 条有区分度的弱线索；", None),
        ("   - 待验证（有线索）：只有 1 条弱线索，或简历写「备考/待考/计划考」；", None),
        ("   - 待验证（未提及）：简历完全没提；没写不等于不会，仍计入命中率分母，面试时逐项确认。", None),
        ("4. 「判据方式」：简历证据 = 看简历文字即可判定；面试核实 = 简历没写也不算不匹配，", None),
        ("   不计入分母（这批标签暂时都是「简历证据」）。", None),
        ("5. 「人工修正」列写你的意见即可，例如：", None),
        ("   - 「删除：标本」（这个词太泛，几乎人人命中）", None),
        ("   - 「补充：质谱仪、流式细胞仪」", None),
        ("   - 「这条不算经验：轮转科室不算」", None),
        ("6. 如果你认为某个标签本身定义不对（例如应拆成两个标签，或该归到经验维度），", None),
        ("   直接写在「人工修正」里，不用改词。填好后把文件发我，我按你的修正改词表并重跑一致率。", None),
    ]
    for index, (text, font) in enumerate(lines, start=1):
        cell = ws.cell(row=index, column=1, value=text)
        cell.alignment = Alignment(wrap_text=True, vertical="top")
        if font:
            cell.font = font
        if text.startswith("   -"):
            cell.alignment = Alignment(wrap_text=True, vertical="top", indent=1)


def build(audit_path: Path, out_path: Path) -> dict:
    audit = _load(audit_path)
    resume_files = [row["file"] for row in audit.get("rows") or []]
    tags = _tag_rows(audit)
    wb = Workbook()
    spec = wb.active
    spec.title = "标签口径"
    _write_spec_sheet(spec, tags, resume_files)
    _write_notes_sheet(wb.create_sheet("填表说明"))
    out_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out_path)
    return {"out": str(out_path), "tags": len(tags), "resumes": len(resume_files)}


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--audit", required=True, help="audit.json produced by eval-matching.py")
    parser.add_argument("--out", required=True, help="xlsx path to write")
    args = parser.parse_args()
    result = build(Path(args.audit), Path(args.out))
    print(json.dumps(result, ensure_ascii=False))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
