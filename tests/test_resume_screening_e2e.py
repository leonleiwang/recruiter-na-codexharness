"""端到端回归：真实文件（DOCX / PDF / 图片）→ service → 匹配报告。

文本输入之外的格式由 ingest 层负责；本文件用现场生成的真实文件验证整条链路
（字段抽取 → 三维标签 → 匹配），并确认图片走 Windows OCR 也能得到同样的结论。
"""
from __future__ import annotations

import shutil
import sys
import tempfile
from pathlib import Path

import pytest

from rn.resume_screening import service

RESUME_LINES = [
    "张晓明",
    "电话：13812345678    邮箱：zhangxm@example.com",
    "性别：男    年龄：27    现居地：杭州市西湖区",
    "学历：本科    专业：医学检验技术    毕业院校：浙江中医药大学    毕业年份：2021",
    "工作年限：3年    期望城市：杭州    求职意向：检验技术员",
    "证书：检验师资格证    技能：PCR、生化分析仪、室内质控",
    "工作经历",
    "2021.07-2024.08    杭州迪安医学检验所    检验技术员",
    "负责临检标本接收、上机检测、室内质控与结果审核，带教2名实习生，体检季日样本量800+。",
    "自我评价",
    "工作细心严谨，能承受体检季高强度加班，主动复盘异常案例，责任心强。",
]

JD_TEXT = """岗位职责：
1. 负责检验标本的接收、上机检测、结果审核与报告发布，执行室内质控与室间质评。
任职要求：
1. 大专及以上学历，医学检验技术/临床医学/生物技术等相关专业；
2. 1年以上医院检验科或第三方医学检验所检验工作经历；
3. 持有检验士/检验师资格证；
4. 熟悉生化分析仪、PCR等仪器的操作；
5. 细心严谨、责任心强，能承受高峰期工作压力。
"""


@pytest.fixture()
def workdir():
    base = Path(tempfile.mkdtemp(prefix="rn-e2e-"))
    try:
        yield base
    finally:
        shutil.rmtree(base, ignore_errors=True)


def _make_docx(path: Path) -> None:
    from docx import Document

    document = Document()
    document.add_heading("个人简历", level=1)
    for line in RESUME_LINES:
        document.add_paragraph(line)
    document.save(path)


def _make_pdf(path: Path) -> None:
    import fitz

    document = fitz.open()
    page = document.new_page(width=595, height=842)
    y = 60
    page.insert_text((50, y), "个人简历", fontsize=18, fontname="china-s")
    y += 32
    for line in RESUME_LINES:
        page.insert_text((50, y), line, fontsize=10, fontname="china-s")
        y += 22
    document.save(path)
    document.close()


def _make_png(path: Path) -> None:
    from PIL import Image, ImageDraw, ImageFont

    image = Image.new("RGB", (1000, 700), "white")
    draw = ImageDraw.Draw(image)
    font = ImageFont.truetype(r"C:\Windows\Fonts\msyh.ttc", 22)
    y = 40
    for line in RESUME_LINES:
        draw.text((36, y), line, font=font, fill="black")
        y += 44
    image.save(path)


def _assert_report(report: dict, *, expect_engine: str) -> None:
    assert report["fields"]["name"] == "张晓明"
    assert report["fields"]["education"]["level"] == "本科"
    assert "检验师资格证" in report["fields"]["certifications"]
    assert report["fields"]["work_experiences"], "工作经历应被解析"
    assert report["matching"]["overall_percent"] > 50.0
    assert report["matching"]["label"] in ("高匹配", "较高匹配")
    assert expect_engine in report["document"]["engines"]


def test_docx_resume_end_to_end(workdir: Path):
    path = workdir / "resume.docx"
    _make_docx(path)
    report = service.analyze_resume_file(path, job_text=JD_TEXT)
    _assert_report(report, expect_engine="python-docx")
    assert report["document"]["media_type"] == "docx"


def test_pdf_resume_end_to_end(workdir: Path):
    path = workdir / "resume.pdf"
    _make_pdf(path)
    report = service.analyze_resume_file(path, job_text=JD_TEXT)
    _assert_report(report, expect_engine="pymupdf")
    assert report["document"]["media_type"] == "pdf"


@pytest.mark.skipif(
    sys.platform != "win32" or not shutil.which("powershell"),
    reason="Windows OCR 只在 Windows + powershell.exe 可用",
)
def test_image_resume_end_to_end_via_windows_ocr(workdir: Path):
    path = workdir / "resume.png"
    _make_png(path)
    report = service.analyze_resume_file(path, job_text=JD_TEXT)
    _assert_report(report, expect_engine="windows-ocr")
    assert report["document"]["media_type"] == "image"


@pytest.mark.skipif(
    sys.platform != "win32" or not shutil.which("powershell"),
    reason="Windows OCR 只在 Windows + powershell.exe 可用",
)
def test_bytes_upload_image_matches_file_analysis(workdir: Path):
    path = workdir / "resume.png"
    _make_png(path)
    report = service.analyze_resume_bytes(path.read_bytes(), filename="upload.png", job_text=JD_TEXT)
    assert report["fields"]["name"] == "张晓明"
    assert report["matching"]["overall_percent"] > 50.0
