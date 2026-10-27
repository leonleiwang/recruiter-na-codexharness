"""回归测试：代码审查 19 项修复（requirements / matching / fields / ingest / CLI / 写盘）。

这些用例对应修复前的真实缺陷，例如：
- 敏感属性（性别/年龄/婚育）曾进入岗位要求与匹配排序；
- 只有硬性条件的 JD 曾被误判成「岗位要求为空，无法评估」；
- 职责句（「保证设备稳定运行」）曾被当成特质证据；
- 证书否定语境（「尚未取得检验师资格证」）曾被判成满足；
- 专业曾用整篇简历做子串匹配（文秘专业被「医学检验所」工作经历命中）；
- 只有年份的工作区间（「2019 - 2022」）曾被算成 4 年；
- 混合页 PDF 的图片页曾被跳过、DOCX 文本框文字曾丢失；
- CLI 曾无法读取 GBK 的 JD 文件、退出码不区分输入错误与分析失败；
- 报告曾静默覆盖旧结果，且 stem 可控时可能路径逃逸。
"""
from __future__ import annotations

import io
import json
from pathlib import Path

import pytest

from rn.resume_screening import cli, service
from rn.resume_screening.analysis import fields, matching, requirements, tags
from rn.resume_screening.analysis.requirements import RequirementsUnavailableError
from rn.resume_screening.ingest.document import load_bytes

STRONG_FIELDS = {
    "name": "张晓明",
    "location": "杭州市西湖区",
    "expected_city": "杭州",
    "education": {"level": "本科", "school": "浙江中医药大学", "major": "医学检验技术"},
    "certifications": ["检验师资格证"],
    "skills": ["PCR", "生化分析仪"],
    "raw_text": (
        "张晓明 男 27岁 本科 医学检验技术 浙江中医药大学\n"
        "工作经历：迪安诊断 检验技术员 2021.07-2024.08\n"
        "负责临检标本上机检测、室内质控与结果审核，熟悉 SOP 与标准操作规程。\n"
        "证书：检验师资格证\n"
    ),
}

MULTI_HARD_JD = "大专及以上学历，医学检验技术专业，持有检验士资格证；工作地点：杭州市西湖区。"


def _match(jd_text: str, candidate: dict) -> dict:
    jd = requirements.requirements_from_jd_text(jd_text)
    return matching.match_requirements(jd, tags.extract_resume_tags(candidate), candidate)


def _hints(report: dict) -> dict[str, dict]:
    return {hint["filter"]: hint for hint in report.get("hard_filter_hints") or []}


def _png_bytes(color=(255, 255, 255), size=(120, 60)) -> bytes:
    from PIL import Image

    buf = io.BytesIO()
    Image.new("RGB", size, color).save(buf, format="PNG")
    return buf.getvalue()


class _FakeOcrEngine:
    """满足 recognize(png_bytes) -> str 的最小测试引擎（不依赖本机 OCR）。"""

    name = "fake-ocr"

    def recognize(self, png: bytes) -> str:  # noqa: ARG002 - 接口要求
        return "识别文本：熟悉 PCR 与生化分析仪"


# ---------------------------------------------------------------------------
# P0：敏感属性不得进入岗位要求 / 排序
# ---------------------------------------------------------------------------


def test_sensitive_attributes_excluded_from_requirements():
    doc = requirements.requirements_from_jd_text(
        "1. 年龄35岁以下，限男性；2. 已婚已育者优先；3. 能独立完成上机检测与结果审核。"
    )
    names = [row["name"] for row in doc["requirements"]]
    assert names, "正常要求不能被敏感行一起丢掉"
    assert all("年龄" not in name and "婚" not in name and "男" not in name for name in names)
    assert len(doc["sensitive_excluded"]) == 2
    assert any("敏感属性" in warning for warning in doc["warnings"])


def test_sensitive_fragment_is_split_out_of_mixed_line():
    # 同一句里「年龄35岁以下」与正常要求并存时，只剔除敏感片段，正常要求保留。
    doc = requirements.requirements_from_jd_text("年龄35岁以下；能独立完成上机检测")
    assert doc["sensitive_excluded"] == ["年龄35岁以下"]
    assert [row["name"] for row in doc["requirements"]] == ["规范操作与执行能力"]


# ---------------------------------------------------------------------------
# P1：未知岗位 / 空要求 / 硬性条件
# ---------------------------------------------------------------------------


def test_unknown_job_without_jd_is_not_fake_requirement():
    doc = requirements.build_job_requirements(job="生物信息工程师")
    assert doc["source"] == "unknown_job"
    assert doc["requirements"] == []
    assert any("--job-text" in warning or "--jd-file" in warning for warning in doc["warnings"])

    report = matching.match_requirements(doc, tags.extract_resume_tags(STRONG_FIELDS), STRONG_FIELDS)
    assert report["requirements_available"] is False
    assert report["label"] == "无法评估（岗位要求为空）"
    assert report["overall_percent"] is None

    with pytest.raises(RequirementsUnavailableError):
        requirements.build_job_requirements(job="生物信息工程师", strict=True)


def test_one_line_multiple_hard_filters_are_split_by_kind():
    doc = requirements.requirements_from_jd_text(MULTI_HARD_JD)
    hard = doc["hard_filters"]
    assert {"education", "certification", "location"} <= set(hard)
    assert hard["location"] == "杭州市西湖区"
    # 地点标签不能混进学历/证书条件里
    assert "杭州市西湖区" not in hard["education"]


def test_location_label_stripped_and_city_matched():
    report = _match(MULTI_HARD_JD, STRONG_FIELDS)
    hints = _hints(report)
    assert hints["education"]["status"] == "满足"
    assert hints["certification"]["status"] == "满足"
    assert hints["location"]["status"] == "满足"
    # 提示里比对的是城市粒度（「命中 杭州」），不能把「工作地点：」标签混进来
    assert "杭州" in hints["location"]["note"]
    assert "工作地点" not in hints["location"]["requirement"]


def test_jd_with_only_hard_filters_is_still_evaluable():
    # 修复前：只有硬性条件的 JD → requirements 为空 → 报告显示「无法评估」，硬性提示却已判「满足」。
    doc = requirements.requirements_from_jd_text(MULTI_HARD_JD)
    assert doc["requirements"] == []
    assert doc["hard_filters"]
    assert requirements.has_usable_requirements(doc)

    report = _match(MULTI_HARD_JD, STRONG_FIELDS)
    assert report["requirements_available"] is True
    assert report["hard_filters_available"] is True
    assert report["label"] == "仅硬性条件（无三维度要求）"
    assert report["overall_percent"] is None  # 覆盖度不适用，不能伪装成 0%
    assert report["hard_filter_hints"]
    assert any("硬性条件" in note for note in report["notes"])

    # strict 模式也不应把「只有硬性条件」当成空要求拒绝
    built = requirements.build_job_requirements(job_text=MULTI_HARD_JD, strict=True)
    assert built["hard_filters"]


# ---------------------------------------------------------------------------
# P1：职责句 / 证书 / 专业
# ---------------------------------------------------------------------------


def test_duty_sentence_is_not_trait_evidence():
    doc = requirements.requirements_from_jd_text(
        "负责仪器设备的日常保养，保证设备稳定运行；能承受较强工作压力，性格开朗。"
    )
    dims = [(row["dimension"], row["name"]) for row in doc["requirements"]]
    assert ("trait", "抗压") in dims
    # 「保证设备稳定运行」是职责描述，不能生成「稳定/严谨」类特质要求
    assert not any(dimension == "trait" and "稳定" in name for dimension, name in dims)


def test_cert_negation_needs_verification():
    candidate = dict(STRONG_FIELDS)
    candidate["certifications"] = []
    candidate["raw_text"] = (
        "张晓明 本科 文秘 某大学\n"
        "工作经历：迪安诊断 检验技术员 2021.07-2024.08\n"
        "已报名考试，尚未取得检验师资格证。\n"
    )
    report = _match("持有检验士资格证", candidate)
    hint = _hints(report)["certification"]
    assert hint["status"] == "待核实"
    assert "未取得" in hint["note"] or "待考" in hint["note"]


def test_unrelated_certificate_does_not_satisfy_检验_requirement():
    candidate = dict(STRONG_FIELDS)
    candidate["certifications"] = ["教师资格证"]
    candidate["raw_text"] = "证书：教师资格证"
    report = _match("持有检验士资格证", candidate)
    hint = _hints(report)["certification"]
    assert hint["status"] == "待核实"  # 修复前「资格证」子串会误判满足


def test_major_only_matched_within_education_scope():
    candidate = dict(STRONG_FIELDS)
    candidate["education"] = {"level": "本科", "school": "某大学", "major": "文秘"}
    candidate["raw_text"] = (
        "张晓明 本科 文秘 某大学\n"
        "工作经历：某医学检验所 检验技术员 2021.07-2024.08\n"
        "证书：检验师资格证\n"
    )
    report = _match("医学检验技术专业，本科及以上", candidate)
    hint = _hints(report)["education"]
    assert hint["status"] == "待核实"  # 工作经历里出现「医学检验所」不算专业匹配
    assert "专业未确认" in hint["note"]


# ---------------------------------------------------------------------------
# P1：fields 抽取细节
# ---------------------------------------------------------------------------


def test_year_only_work_range_is_not_overestimated():
    result = fields.extract_resume_fields("工作经历\n2019 - 2022 迪安诊断 检验员")
    entry = result["work_experiences"][0]
    assert entry["months"] == 36
    # 修复前：区间被补成 2019-01 ~ 2022-12（48 个月）→ 4 年
    assert result["total_work_years"] == 3


def test_skills_are_not_swallowed_by_certification_line():
    result = fields.extract_resume_fields("证书：检验师资格证\n技能：PCR、生化分析仪、室内质控")
    skills = result["skills"]
    assert "PCR" in skills
    assert "生化分析仪" in skills
    assert "室内质控" in skills


def test_reversed_work_range_is_ignored():
    result = fields.extract_resume_fields(
        "工作经历\n2019.07-2021.06 A公司 检验员\n2024.08-2021.07 B公司 检验员"
    )
    assert all(entry["months"] >= 0 for entry in result["work_experiences"])
    assert result["total_work_years"] == 2  # 只算正常的 24 个月，逆序区间不倒扣


def test_spaced_ocr_text_still_produces_capability_tags():
    fields_result = fields.extract_resume_fields("熟 悉 P C R 与 生 化 分 析 仪")
    tags_doc = tags.extract_resume_tags(fields_result)
    capability = tags_doc["tags"]["capability"]
    assert capability, "OCR 字符间空格不应导致三维标签全空"
    assert any(row["name"] == "仪器操作能力" for row in capability)


# ---------------------------------------------------------------------------
# P2：ingest（混合页 PDF / DOCX 文本框 / 编码 / 伪装文本）
# ---------------------------------------------------------------------------


def test_mixed_pdf_page_falls_back_to_ocr():
    import fitz

    pdf = fitz.open()
    page = pdf.new_page()
    # 一页里既有少量矢量文字（< MIXED_PDF_TEXT_CHARS）又有大图：必须按扫描页走 OCR
    page.insert_text((72, 72), "张晓明 13812345678", fontsize=11)
    page.insert_image(fitz.Rect(60, 100, 360, 300), stream=_png_bytes())
    data = pdf.tobytes()
    pdf.close()

    document = load_bytes(data, "mixed.pdf", ocr_engine=_FakeOcrEngine())
    assert "识别文本" in document.text
    assert "fake-ocr" in document.engines
    assert document.meta["pages"][0]["ocr"] is True


def test_docx_textbox_text_is_extracted(tmp_path: Path):
    from docx import Document
    from docx.oxml import parse_xml
    from docx.oxml.ns import nsdecls

    doc = Document()
    doc.add_paragraph("普通段落：张晓明")
    paragraph = doc.add_paragraph()
    run = paragraph.add_run()
    run._r.append(parse_xml(
        f"<w:pict {nsdecls('w')}>"
        "<v:shape xmlns:v='urn:schemas-microsoft-com:vml' style='width:200pt;height:50pt'>"
        "<v:textbox><w:txbxContent><w:p><w:r><w:t>文本框内容：医学检验技术</w:t></w:r></w:p>"
        "</w:txbxContent></v:textbox></v:shape></w:pict>"
    ))
    path = tmp_path / "textbox.docx"
    doc.save(path)

    document = load_bytes(path.read_bytes(), "textbox.docx")
    assert "普通段落" in document.text
    assert "文本框内容" in document.text  # 修复前 python-docx 只读 body 段落，会丢掉文本框


def test_utf16_text_without_bom_is_decoded():
    data = "姓名：张三\n电话：13812345678".encode("utf-16-le")
    document = load_bytes(data, "utf16.txt")
    assert "姓名：张三" in document.text
    assert document.meta["encoding"].startswith("utf-16")


def test_png_with_txt_suffix_is_not_decoded_as_text():
    # PNG 改名成 .txt：不能让二进制被当 UTF-8 强制解码成乱码要求。
    document = load_bytes(_png_bytes(), "image.txt")
    assert document.media_type == "image"
    assert "PNG" not in document.text


# ---------------------------------------------------------------------------
# 打分口径：特质不计分（只展示、只做面试关注点）
# ---------------------------------------------------------------------------


def test_trait_dimension_is_reference_only():
    req = requirements.build_job_requirements(job="检验技术员")
    report = matching.match_requirements(req, tags.extract_resume_tags(STRONG_FIELDS), STRONG_FIELDS)
    dims = report["dimensions"]
    assert dims["capability"]["scored"] is True
    assert dims["experience"]["scored"] is True
    assert dims["trait"]["scored"] is False
    assert set(report["dimension_weights"]) == {"capability", "experience"}
    assert report["scored_dimensions"] == ["capability", "experience"]
    assert report["reference_dimensions"] == ["trait"]

    # 综合覆盖度 = 能力/经验按 0.40/0.35 归一化加权，特质不参与
    manual = (
        dims["capability"]["coverage"] * 0.40 + dims["experience"]["coverage"] * 0.35
    ) / 0.75
    assert report["overall_coverage"] == pytest.approx(manual, abs=0.002)

    # 特质必须项未命中不进缺口、也不进「已证实的必须项」（后者只用于排序）
    assert all(gap["dimension"] != "trait" for gap in report["gaps"])
    scoring_names = {
        item["name"]
        for dim in ("capability", "experience")
        for item in dims[dim]["items"]
    }
    assert set(report["matched_must_have"]) <= scoring_names


def test_trait_only_requirements_have_no_score():
    req = requirements.requirements_from_jd_text("性格开朗，抗压能力强，工作积极主动。")
    assert {row["dimension"] for row in req["requirements"]} == {"trait"}
    report = matching.match_requirements(req, tags.extract_resume_tags(STRONG_FIELDS), STRONG_FIELDS)
    assert report["overall_percent"] is None
    assert report["label"] == "仅参考信息（无计分要求）"
    assert report["dimension_weights"] == {}


def test_report_marks_trait_as_not_scored():
    req = requirements.build_job_requirements(job="检验技术员")
    report = matching.match_requirements(req, tags.extract_resume_tags(STRONG_FIELDS), STRONG_FIELDS)
    markdown = matching.render_matching_markdown(report)
    assert "仅能力/经验计分" in markdown
    assert "特质（不计分，仅参考）" in markdown


def test_trait_tags_do_not_change_overall_percent():
    import copy

    req = requirements.build_job_requirements(job="检验技术员")
    with_traits = dict(STRONG_FIELDS)
    with_traits["self_evaluation"] = "工作细心严谨，抗压能力强，责任心强，积极主动学习。"
    with_traits["raw_text"] = STRONG_FIELDS["raw_text"] + "\n自我评价：工作细心严谨，能承受高强度加班，主动复盘异常。"

    tags_doc = tags.extract_resume_tags(with_traits)
    stripped = copy.deepcopy(tags_doc)
    stripped["tags"]["trait"] = []  # 同一份简历，只去掉特质标签

    full_report = matching.match_requirements(req, tags_doc, with_traits)
    stripped_report = matching.match_requirements(req, stripped, with_traits)
    # 特质标签数量变化不影响综合匹配度（特质不计分）
    assert stripped_report["overall_coverage"] == full_report["overall_coverage"]
    assert stripped_report["overall_percent"] == full_report["overall_percent"]


# ---------------------------------------------------------------------------
# P2：CLI 退出码 / 编码 / 写盘安全
# ---------------------------------------------------------------------------


def test_cli_reads_gbk_jd_file_and_writes_nothing(tmp_path: Path):
    jd_path = tmp_path / "jd_gbk.txt"
    jd_path.write_bytes("大专及以上学历，医学检验技术专业；持有检验士资格证。".encode("gb18030"))
    resume_path = tmp_path / "resume.txt"
    resume_path.write_text("张晓明 本科 医学检验技术 检验师资格证 13812345678", encoding="utf-8")

    rc = cli.main(["analyze", str(resume_path), "--jd-file", str(jd_path), "--no-write", "--json"])
    assert rc == cli.EXIT_OK  # 修复前 GBK JD 文件直接 UnicodeDecodeError


def test_cli_rejects_text_and_files_together():
    assert cli.main(["analyze", "a.txt", "--text", "x"]) == cli.EXIT_INPUT_ERROR


def test_cli_rejects_empty_requirements_by_default(tmp_path: Path):
    rc = cli.main(["analyze", "--text", "张三", "--job", "生物信息工程师", "--no-write"])
    assert rc == cli.EXIT_INPUT_ERROR
    rc_allowed = cli.main(
        ["analyze", "--text", "张三", "--job", "生物信息工程师", "--no-write", "--allow-empty-requirements"]
    )
    assert rc_allowed == cli.EXIT_OK


def test_cli_batch_all_failed_returns_analysis_error(tmp_path: Path):
    jd_path = tmp_path / "jd.txt"
    jd_path.write_text(MULTI_HARD_JD, encoding="utf-8")
    rc = cli.main([
        "analyze", str(tmp_path / "a.pdf"), str(tmp_path / "b.docx"),
        "--jd-file", str(jd_path), "--out", str(tmp_path / "out"),
    ])
    assert rc == cli.EXIT_ANALYSIS_ERROR  # 修复前全部失败也返回 0
    assert not (tmp_path / "out").exists()


def test_write_report_refuses_overwrite_and_escapes_stem(tmp_path: Path):
    report = service.analyze_text("张晓明 本科 医学检验技术", job_text=MULTI_HARD_JD)
    out_dir = tmp_path / "runs"
    paths = service.write_report(report, out_dir=out_dir, stem="../../evil")
    assert Path(paths["json"]).resolve() == (out_dir / "evil.json").resolve()
    assert Path(paths["markdown"]).resolve() == (out_dir / "evil.md").resolve()

    with pytest.raises(FileExistsError):
        service.write_report(report, out_dir=out_dir, stem="../../evil")

    paths_again = service.write_report(report, out_dir=out_dir, stem="../../evil", overwrite=True)
    assert Path(paths_again["json"]).exists()


def test_write_batch_report_volume_is_json_serializable(tmp_path: Path):
    batch = service.analyze_resume_files(
        [], job_text=MULTI_HARD_JD, strict_requirements=False,
    )
    paths = service.write_batch_report(batch, out_dir=tmp_path / "batch")
    payload = json.loads(Path(paths["json"]).read_text(encoding="utf-8"))
    assert payload["reports"] == []
    assert "batch" in payload["schema_version"] or "analysis" in payload["schema_version"]
